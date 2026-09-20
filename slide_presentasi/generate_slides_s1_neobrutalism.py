"""
=====================================================================
GENERATOR SLIDE SESI 01 (EDISI MASTERPIECE 18 SLIDE LENGKAP)
Mata Kuliah: Pemrograman Berbasis Perangkat Bergerak (STSI4303 / MSIM4401)
Institusi  : Universitas Terbuka — Fakultas Sains dan Teknologi (FST)
Dosen      : Anton Prafanto, S.Kom., M.T.
Standar    : Neo-Brutalism Murni (16:9, Zero Border Radius, Bold Borders)
=====================================================================
"""

import os
import sys

# Pastikan output konsol mendukung UTF-8 di Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# --- PALET WARNA NEO-BRUTALISM ---
COLOR_BG = RGBColor(250, 248, 245)        # #FAF8F5 (Warm Cream Canvas)
COLOR_BLACK = RGBColor(0, 0, 0)           # #000000 (Pure Black Border & Hard Shadows)
COLOR_WHITE = RGBColor(255, 255, 255)     # #FFFFFF (Pure White Card)
COLOR_YELLOW = RGBColor(255, 230, 0)      # #FFE600 (Canary Yellow)
COLOR_CYAN = RGBColor(56, 189, 248)       # #38BDF8 (Electric Cyan)
COLOR_MINT = RGBColor(74, 222, 128)       # #4ADE80 (Neo Mint / Green)
COLOR_CORAL = RGBColor(251, 113, 133)     # #FB7185 (Neo Coral / Rose)
COLOR_PURPLE = RGBColor(192, 132, 252)    # #C084FC (Lavender Purple)
COLOR_CODE_BG = RGBColor(18, 18, 18)      # #121212 (Deep Black Code Box)
COLOR_CODE_TEXT = RGBColor(245, 245, 245) # #F5F5F5 (Crisp Monospace)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_CODE = "Consolas"

SHADOW_OFFSET = Inches(0.08)
GITHUB_BASE_URL = "https://github.com/antonprafanto/tuweb_mobile2025/blob/main/contoh_kode_program/sesi_01_lingkungan_dan_tools"

class NeoBrutalistDeckBuilder:
    def __init__(self, course_name="STSI4303 • PEMROGRAMAN BERBASIS PERANGKAT BERGERAK"):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]
        self.course_name = course_name

    def _set_canvas_bg(self, slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def _add_neobrutal_card(self, slide, x, y, w, h, fill_color=COLOR_WHITE, has_shadow=True, border_width=Pt(2.5)):
        if has_shadow:
            shadow = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, x + SHADOW_OFFSET, y + SHADOW_OFFSET, w, h
            )
            shadow.fill.solid()
            shadow.fill.fore_color.rgb = COLOR_BLACK
            shadow.line.fill.background()

        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, x, y, w, h
        )
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = COLOR_BLACK
        card.line.width = border_width
        return card

    def _add_header(self, slide, tag, title, tag_color=COLOR_YELLOW):
        tag_w = Inches(len(tag) * 0.125 + 0.6)
        self._add_neobrutal_card(slide, Inches(0.9), Inches(0.4), tag_w, Inches(0.4), fill_color=tag_color, has_shadow=True)
        tb_tag = slide.shapes.add_textbox(Inches(0.9), Inches(0.4), tag_w, Inches(0.4))
        p_tag = tb_tag.text_frame.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.name = FONT_HEADING
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_BLACK

        tb_title = slide.shapes.add_textbox(Inches(0.9), Inches(0.88), Inches(11.5), Inches(0.8))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_BLACK

    # 1. Slide Cover
    def add_cover(self, meeting_num, title, subtitle, presenter_info):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_canvas_bg(slide)

        self._add_neobrutal_card(slide, Inches(0.9), Inches(0.9), Inches(11.533), Inches(5.7), fill_color=COLOR_WHITE)

        tb = slide.shapes.add_textbox(Inches(1.4), Inches(1.3), Inches(10.5), Inches(4.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"✦ {self.course_name.upper()} • SESI {meeting_num:02d} ✦"
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_BLACK
        p0.space_after = Pt(16)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_BLACK
        p1.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = FONT_BODY
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(60, 60, 60)
        p2.space_after = Pt(36)

        p3 = tf.add_paragraph()
        p3.text = f"■  {presenter_info}  ■  FAKULTAS SAINS DAN TEKNOLOGI • UNIVERSITAS TERBUKA"
        p3.font.name = FONT_HEADING
        p3.font.size = Pt(11.5)
        p3.font.bold = True
        p3.font.color.rgb = COLOR_BLACK

    # 2. Slide Split Screen: Konsep & Tips di Kiri + Kode / Terminal di Kanan
    def add_concept_with_code(self, tag, title, bullets, code_snippet, filename="terminal_commands.sh", tip=None, tag_color=COLOR_YELLOW, full_code_file=None):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_canvas_bg(slide)
        self._add_header(slide, tag, title, tag_color=tag_color)

        # Left Explanations Card (4.9 Inches)
        self._add_neobrutal_card(slide, Inches(0.9), Inches(1.8), Inches(4.9), Inches(5.1), fill_color=COLOR_WHITE)

        tb_l = slide.shapes.add_textbox(Inches(1.15), Inches(1.95), Inches(4.4), Inches(4.8))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True

        p_lt = tf_l.paragraphs[0]
        p_lt.text = "KONSEP & ALUR KERJA"
        p_lt.font.name = FONT_HEADING
        p_lt.font.size = Pt(11.5)
        p_lt.font.bold = True
        p_lt.font.color.rgb = COLOR_BLACK
        p_lt.space_after = Pt(8)

        for b in bullets:
            p = tf_l.add_paragraph()
            p.text = f"■  {b}"
            p.font.name = FONT_BODY
            p.font.size = Pt(10.5)
            p.font.color.rgb = RGBColor(20, 20, 20)
            p.space_after = Pt(5)

        if tip:
            p_tip = tf_l.add_paragraph()
            p_tip.text = f"💡 Tips Praktis: {tip}"
            p_tip.font.name = FONT_HEADING
            p_tip.font.size = Pt(10)
            p_tip.font.bold = True
            p_tip.font.color.rgb = COLOR_BLACK

        # Right Sharp Black Code Box (6.433 Inches)
        self._add_neobrutal_card(slide, Inches(6.0), Inches(1.8), Inches(6.433), Inches(5.1), fill_color=COLOR_CODE_BG)

        # Code Header Strip
        c_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.0), Inches(1.8), Inches(6.433), Inches(0.48))
        c_bar.fill.solid()
        c_bar.fill.fore_color.rgb = tag_color
        c_bar.line.color.rgb = COLOR_BLACK
        c_bar.line.width = Pt(2.5)

        tb_cb = slide.shapes.add_textbox(Inches(6.2), Inches(1.86), Inches(6.0), Inches(0.38))
        p_cb = tb_cb.text_frame.paragraphs[0]
        p_cb.text = f"{filename}  |  KODE IMPLEMENTASI"
        p_cb.font.name = FONT_CODE
        p_cb.font.size = Pt(10.5)
        p_cb.font.bold = True
        p_cb.font.color.rgb = COLOR_BLACK

        code_box_h = Inches(3.68) if full_code_file else Inches(4.35)
        tb_r = slide.shapes.add_textbox(Inches(6.2), Inches(2.32), Inches(6.0), code_box_h)
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p_code = tf_r.paragraphs[0]
        p_code.text = code_snippet
        p_code.font.name = FONT_CODE
        p_code.font.size = Pt(9.0)
        p_code.font.color.rgb = COLOR_CODE_TEXT

        # Neo-Brutalist Runnable Code Hyperlink Button
        if full_code_file:
            btn_x = Inches(6.15)
            btn_y = Inches(6.16)
            btn_w = Inches(6.133)
            btn_h = Inches(0.60)

            btn_shadow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x + Inches(0.04), btn_y + Inches(0.04), btn_w, btn_h)
            btn_shadow.fill.solid()
            btn_shadow.fill.fore_color.rgb = COLOR_BLACK
            btn_shadow.line.fill.background()

            btn = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x, btn_y, btn_w, btn_h)
            btn.fill.solid()
            btn.fill.fore_color.rgb = COLOR_YELLOW
            btn.line.color.rgb = COLOR_BLACK
            btn.line.width = Pt(2.0)

            tb_btn = slide.shapes.add_textbox(btn_x + Inches(0.08), btn_y + Inches(0.03), btn_w - Inches(0.16), btn_h - Inches(0.06))
            tf_btn = tb_btn.text_frame
            tf_btn.word_wrap = True

            p_btn = tf_btn.paragraphs[0]
            run_btn = p_btn.add_run()
            run_btn.text = f"▶ BUKA KODE LENGKAP DI GITHUB ({full_code_file})"
            run_btn.font.name = FONT_HEADING
            run_btn.font.size = Pt(9.5)
            run_btn.font.bold = True
            run_btn.font.color.rgb = COLOR_BLACK
            run_btn.hyperlink.address = f"{GITHUB_BASE_URL}/{full_code_file}"

            p_btn_sub = tf_btn.add_paragraph()
            p_btn_sub.text = "💡 Klik untuk buka di browser / salin source code langsung"
            p_btn_sub.font.name = FONT_BODY
            p_btn_sub.font.size = Pt(7.5)
            p_btn_sub.font.color.rgb = RGBColor(60, 60, 60)

    # 3. Slide Lab Quest
    def add_lab_quest(self, meeting_num, title, goals, success_criteria, full_code_file="slide_17_lab_quest_01_profil_mahasiswa.html"):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_canvas_bg(slide)
        self._add_header(slide, "LAB QUEST MANDIRI", f"Sesi {meeting_num:02d}: {title}", tag_color=COLOR_YELLOW)

        self._add_neobrutal_card(slide, Inches(0.9), Inches(1.8), Inches(11.533), Inches(5.1), fill_color=COLOR_WHITE)

        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.8), Inches(11.533), Inches(0.52))
        banner.fill.solid()
        banner.fill.fore_color.rgb = COLOR_YELLOW
        banner.line.color.rgb = COLOR_BLACK
        banner.line.width = Pt(2.5)

        tb_b = slide.shapes.add_textbox(Inches(1.2), Inches(1.86), Inches(10.5), Inches(0.4))
        p_b = tb_b.text_frame.paragraphs[0]
        p_b.text = "🎯 TANTANGAN PRAKTIKUM MANDIRI BERORIENTASI HASIL (OUTCOME-BASED)"
        p_b.font.name = FONT_HEADING
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_BLACK

        # Kolom Kiri: Sasaran Proyek
        self._add_neobrutal_card(slide, Inches(1.2), Inches(2.55), Inches(5.2), Inches(4.0), fill_color=RGBColor(248, 250, 252))
        tb_g = slide.shapes.add_textbox(Inches(1.35), Inches(2.7), Inches(4.9), Inches(3.7))
        tf_g = tb_g.text_frame
        tf_g.word_wrap = True
        p_gt = tf_g.paragraphs[0]
        p_gt.text = "SASARAN & SKENARIO KASUS"
        p_gt.font.name = FONT_HEADING
        p_gt.font.size = Pt(11.5)
        p_gt.font.bold = True
        p_gt.font.color.rgb = COLOR_BLACK
        p_gt.space_after = Pt(8)

        for g in goals:
            p = tf_g.add_paragraph()
            p.text = f"✦  {g}"
            p.font.name = FONT_BODY
            p.font.size = Pt(10.5)
            p.font.color.rgb = RGBColor(30, 41, 59)
            p.space_after = Pt(5)

        # Kolom Kanan: Kriteria Sukses & Solusi
        self._add_neobrutal_card(slide, Inches(6.7), Inches(2.55), Inches(5.4), Inches(4.0), fill_color=RGBColor(240, 253, 244))
        tb_s = slide.shapes.add_textbox(Inches(6.85), Inches(2.7), Inches(5.1), Inches(3.7))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        p_st = tf_s.paragraphs[0]
        p_st.text = "KRITERIA SUKSES & SOLUSI MASTER"
        p_st.font.name = FONT_HEADING
        p_st.font.size = Pt(11.5)
        p_st.font.bold = True
        p_st.font.color.rgb = RGBColor(22, 101, 52)
        p_st.space_after = Pt(8)

        for sc in success_criteria:
            p = tf_s.add_paragraph()
            p.text = f"✔  {sc}"
            p.font.name = FONT_BODY
            p.font.size = Pt(10.5)
            p.font.color.rgb = RGBColor(20, 83, 45)
            p.space_after = Pt(5)

        # Tombol Buka Solusi Master
        btn_x = Inches(6.85)
        btn_y = Inches(5.7)
        btn_w = Inches(5.1)
        btn_h = Inches(0.55)

        btn_shadow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x + Inches(0.04), btn_y + Inches(0.04), btn_w, btn_h)
        btn_shadow.fill.solid()
        btn_shadow.fill.fore_color.rgb = COLOR_BLACK
        btn_shadow.line.fill.background()

        btn = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x, btn_y, btn_w, btn_h)
        btn.fill.solid()
        btn.fill.fore_color.rgb = COLOR_YELLOW
        btn.line.color.rgb = COLOR_BLACK
        btn.line.width = Pt(2.0)

        tb_btn = slide.shapes.add_textbox(btn_x + Inches(0.1), btn_y + Inches(0.05), btn_w - Inches(0.2), btn_h - Inches(0.1))
        tf_btn = tb_btn.text_frame
        p_btn = tf_btn.paragraphs[0]
        r_btn = p_btn.add_run()
        r_btn.text = f"▶ BUKA MASTER SOLUSI LAB QUEST ({full_code_file})"
        r_btn.font.name = FONT_HEADING
        r_btn.font.size = Pt(9.5)
        r_btn.font.bold = True
        r_btn.font.color.rgb = COLOR_BLACK
        r_btn.hyperlink.address = f"{GITHUB_BASE_URL}/{full_code_file}"

    def save(self, filepath):
        self.prs.save(filepath)
        print(f"✅ Slide presentasi berhasil dibuat: {filepath}")


def build_sesi_01_deck():
    builder = NeoBrutalistDeckBuilder(course_name="STSI4303 • PEMROGRAMAN BERBASIS PERANGKAT BERGERAK")

    # Slide 01: Cover
    builder.add_cover(
        meeting_num=1,
        title="PENGENALAN LINGKUNGAN PENGEMBANGAN\nAPLIKASI PIRANTI BERGERAK",
        subtitle="Evolusi Mobile Hybrid, Setup Node.js & Git, serta Zero-Friction Lab Lingkungan Kerja",
        presenter_info="Anton Prafanto, S.Kom., M.T. • Program Studi S1 Sistem Informasi / Informatika"
    )

    # Slide 02: Relevansi Industri & Efisiensi Hybrid
    builder.add_concept_with_code(
        tag="RELEVANSI INDUSTRI",
        title="Mengapa Industri & Universitas Terbuka Memilih Hybrid?",
        bullets=[
            "Mahasiswa UT tersebar di seluruh Indonesia dengan ribuan variasi ponsel.",
            "Pendekatan Native murni mewajibkan 2 tim terpisah (Kotlin & Swift).",
            "Pendekatan Hybrid cukup 1 basis kode web standar (HTML, CSS, TS, Vue).",
            "Efisiensi waktu dan biaya mencapai ~66% lebih hemat dan cepat rilis.",
            "Dapat didistribusikan ke Web (PWA), Google Play Store, dan Apple App Store."
        ],
        code_snippet="""// Kalkulasi Penghematan Waktu & Sumber Daya:
const timNative = {
  androidKotlin: '1 Orang Developer',
  iosSwift:      '1 Orang Developer',
  webPortal:     '1 Orang Developer'
}; // Total: 3 Basis Kode Terpisah!

const timHybridIonic = {
  singleCodebase: '1 Orang Menguasai Web Standard'
}; // 1 Kode untuk Android, iOS, & Web PWA!""",
        filename="slide_02_relevansi_mobile_hybrid.html",
        tip="Fokus pada penguasaan Web Standards membuka peluang karir di web maupun mobile!",
        tag_color=COLOR_YELLOW,
        full_code_file="slide_02_relevansi_mobile_hybrid.html"
    )

    # Slide 03: Komparasi 3 Arsitektur
    builder.add_concept_with_code(
        tag="ARSITEKTUR MOBILE",
        title="Komparasi Tiga Arsitektur: Native vs Hybrid vs Cross-Platform",
        bullets=[
            "Native (Kotlin/Swift): Akses hardware langsung, performa grafis 3D tinggi.",
            "Hybrid (Ionic + Capacitor): Dibungkus dalam WebView native berkinerja tinggi.",
            "Cross-Platform (Flutter): Menggambar piksel sendiri via engine grafis C++.",
            "Untuk aplikasi data, akademik, dan formulir, Hybrid adalah raja efisiensi.",
            "Akses hardware dijembatani oleh Capacitor Runtime (Camera, GPS, Storage)."
        ],
        code_snippet="""┌─────────────────────────────────────────────────┐
│        APLIKASI WEB (VUE 3 + TYPESCRIPT)        │
├─────────────────────────────────────────────────┤
│    JEMBATAN RUNTIME (CAPACITOR NATIVE PLUGINS)  │
├─────────────────────────────────────────────────┤
│    SISTEM OPERASI SMARTPHONE (ANDROID / IOS)    │
└─────────────────────────────────────────────────┘
*Satu lapis web yang berdaya guna di semua sistem operasi.""",
        filename="slide_03_komparasi_arsitektur_mobile.html",
        tip="Ionic Framework menghilangkan keharusan belajar bahasa native baru dari awal.",
        tag_color=COLOR_CYAN,
        full_code_file="slide_03_komparasi_arsitektur_mobile.html"
    )

    # Slide 04: Ekosistem Ionic & Web Components CDN
    builder.add_concept_with_code(
        tag="EKOSISTEM IONIC",
        title="Web Components Bawaan Ionic: Desain Mobile Siap Pakai",
        bullets=[
            "Tidak perlu merancang tombol, kartu, atau header dari nol.",
            "Komponen Ionic sudah mengadaptasi standar Material Design (Android).",
            "Dapat dijalankan secara instan via CDN tanpa instalasi apa pun di awal.",
            "Cukup menyisipkan script @ionic/core di berkas HTML.",
            "Sangat ramah bagi mahasiswa dengan laptop spesifikasi RAM 4GB."
        ],
        code_snippet="""<!-- Cukup tautkan CDN resmi Ionic di berkas HTML -->
<script type="module" 
  src="https://cdn.jsdelivr.net/npm/@ionic/core/dist/ionic/ionic.esm.js">
</script>
<link rel="stylesheet" 
  href="https://cdn.jsdelivr.net/npm/@ionic/core/css/ionic.bundle.css" />

<!-- Gunakan komponen siap pakai khas mobile -->
<ion-app>
  <ion-card>
    <ion-card-header>
      <ion-card-title>Portal Mahasiswa UT</ion-card-title>
    </ion-card-header>
    <ion-button expand="block">Sentuh Saya</ion-button>
  </ion-card>
</ion-app>""",
        filename="slide_04_ekosistem_ionic_web_standards.html",
        tip="Klik ganda file HTML mandiri ini untuk langsung melihat tampilan di Google Chrome!",
        tag_color=COLOR_MINT,
        full_code_file="slide_04_ekosistem_ionic_web_standards.html"
    )

    # Slide 05: Panduan Instalasi Tools
    builder.add_concept_with_code(
        tag="PENYIAPAN TOOLS",
        title="Panduan Alat Tempur Pengembang Mobile di Komputer Lokal",
        bullets=[
            "Node.js LTS (v20.x/v22.x): Mesin runtime JavaScript untuk menjalankan tools.",
            "NPM (Node Package Manager): Pengelola paket pustaka otomatis.",
            "Git SCM: Sistem kendali versi untuk mencatat revisi dan mengumpulkan tugas.",
            "VS Code: Editor kode ringan dengan ekstensi resmi Vue & Ionic.",
            "Google Chrome / Edge: Media uji coba antarmuka mobile paling ringan."
        ],
        code_snippet="""# 1. Unduh dan pasang Node.js LTS dari:
# https://nodejs.org/ (Pilih versi Recommended for Most Users)

# 2. Buka Terminal / PowerShell di VS Code lalu ketik:
node -v
npm -v
git --version

# 3. Pasang ekstensi wajib di VS Code:
# • Vue - Official (Syntax highlight Vue 3)
# • Ionic (Snippet komponen Ionic)""",
        filename="slide_05_panduan_instalasi_tools.js",
        tip="Hindari menginstal Node.js versi 'Current/Bleeding Edge', pilih selalu versi 'LTS'.",
        tag_color=COLOR_YELLOW,
        full_code_file="slide_05_panduan_instalasi_tools.js"
    )

    # Slide 06: Diagnostik Kesiapan Sistem
    builder.add_concept_with_code(
        tag="UJI DIAGNOSTIK",
        title="Skrip Diagnostik Otomatis Pemeriksa Kesiapan Sistem & RAM",
        bullets=[
            "Skrip mandiri buatan dosen untuk memeriksa kesiapan komputer Anda.",
            "Mendeteksi versi Node.js, NPM, Git, serta kapasitas memori RAM fisik.",
            "Memberikan rekomendasi jalur praktikum yang aman (Jalur A, B, atau C).",
            "Mencegah laptop freeze karena beban kerja yang berlebihan.",
            "Cukup jalankan satu baris perintah di terminal VS Code."
        ],
        code_snippet="""// Jalankan skrip diagnostik mandiri:
// node slide_06_diagnostik_environment.js

const os = require('os');
const ramGB = (os.totalmem() / (1024 ** 3)).toFixed(2);

console.log(`🖥️  Sistem Operasi : ${os.type()}`);
console.log(`💾 Kapasitas RAM   : ${ramGB} GB`);

if (ramGB < 6) {
  console.log('👉 Jalur A: Web Preview Standalone (Aman RAM 4GB)');
} else {
  console.log('👉 Jalur B: USB Debugging HP Fisik + scrcpy');
}""",
        filename="slide_06_diagnostik_environment.js",
        tip="Jalankan skrip ini sebelum memulai praktikum untuk memastikan tidak ada tools yang kurang.",
        tag_color=COLOR_CYAN,
        full_code_file="slide_06_diagnostik_environment.js"
    )

    # Slide 07: Anatomi Aplikasi Hybrid
    builder.add_concept_with_code(
        tag="ANATOMI HYBRID",
        title="Tritunggal Web: HTML (Tulang), CSS (Baju) & JavaScript (Saraf)",
        bullets=[
            "HTML: Membentuk elemen struktur (tombol, teks judul, kontainer kartu).",
            "CSS: Mengatur dimensi layar ponsel, sudut melengkung, dan warna biru UT.",
            "JavaScript: Menangani sentuhan jari (*touch event*) dan logika bisnis.",
            "Dalam arsitektur hybrid, ketiga elemen ini dibungkus menjadi paket mobile.",
            "Tersedia fitur interaktif untuk mematikan/menyalakan CSS pada file contoh."
        ],
        code_snippet="""<!-- Struktur (Tulang) -->
<div class="card-profil">
  <h3>Budi Santoso</h3>
  <p>NIM: 041234567 • Sistem Informasi</p>
  <button id="btnAksi">Verifikasi KTM</button>
</div>

<!-- Logika Interaksi (Sistem Saraf) -->
<script>
  document.getElementById('btnAksi').onclick = () => {
    alert('KTM Mahasiswa Terverifikasi Aktif!');
  };
</script>""",
        filename="slide_07_anatomi_hybrid_app.html",
        tip="Coba matikan CSS pada file contoh ini untuk melihat pentingnya peran styling mobile.",
        tag_color=COLOR_MINT,
        full_code_file="slide_07_anatomi_hybrid_app.html"
    )

    # Slide 08: Hello Hybrid World
    builder.add_concept_with_code(
        tag="HELLO WORLD",
        title="Hello Hybrid World: Mockup Aplikasi Mobile Pertama Anda",
        bullets=[
            "Aplikasi mandiri pertama dengan bingkai ponsel (*smartphone frame*).",
            "Dilengkapi bar status atas (jam, indikator baterai, sinyal 4G UT).",
            "Membaca spesifikasi resolusi layar secara dinamis menggunakan JavaScript.",
            "Merespon sentuhan tombol dengan umpan balik visual interaktif.",
            "Membuktikan bahwa membuat aplikasi mobile bisa dimulai dari hari pertama!"
        ],
        code_snippet="""// Mendeteksi spesifikasi resolusi layar WebView:
const infoLayar = {
  lebar: window.innerWidth,
  tinggi: window.innerHeight,
  platform: navigator.platform
};

function prosesSentuh() {
  document.getElementById('notif').innerText = 
    '🎉 Event sentuhan (Touch Event) berhasil diproses!';
}""",
        filename="slide_08_hello_hybrid_app.html",
        tip="Tekan F12 lalu klik ikon ponsel di pojok kiri atas DevTools untuk pengalaman terbaik.",
        tag_color=COLOR_YELLOW,
        full_code_file="slide_08_hello_hybrid_app.html"
    )

    # Slide 09: Chrome Device Toolbar Guide
    builder.add_concept_with_code(
        tag="DEVTOOLS EMULATOR",
        title="3 Langkah Menguji Tampilan Ponsel di Google Chrome DevTools",
        bullets=[
            "Langkah 1: Tekan tombol F12 pada keyboard untuk membuka Developer Tools.",
            "Langkah 2: Tekan kombinasi sakti Ctrl + Shift + M (Toggle Device Toolbar).",
            "Langkah 3: Pilih preset smartphone: Pixel 7, iPhone 14 Pro, atau Responsive.",
            "Bisa menguji rotasi layar (Portrait vs Landscape) secara instan.",
            "Konsumsi memori RAM di bawah 150MB, jauh lebih hemat dibanding emulator AVD."
        ],
        code_snippet="""Shortcut Sakti Google Chrome DevTools:
┌───────────────────────────────────────────────┐
│ F12              -> Buka Panel Inspect Elemen │
│ Ctrl + Shift + M -> Toggle Device Toolbar     │
│ Ctrl + Shift + C -> Inspeksi Elemen Tertentu  │
│ Ctrl + R         -> Muat Ulang Halaman        │
└───────────────────────────────────────────────┘
*Simulasi responsif tanpa perlu memasang Android Studio!""",
        filename="slide_09_chrome_device_toolbar_guide.html",
        tip="Gunakan fitur throttling jaringan di DevTools untuk mensimulasikan koneksi 3G pedesaan.",
        tag_color=COLOR_CYAN,
        full_code_file="slide_09_chrome_device_toolbar_guide.html"
    )

    # Slide 10: Viewport Meta Tag
    builder.add_concept_with_code(
        tag="VIEWPORT SCALING",
        title="Peran Krusial Viewport Meta Tag: Skala Presisi 1:1 di Layar HP",
        bullets=[
            "Secara bawaan, peramban smartphone menganggap semua web selebar 980px.",
            "Akibatnya, tanpa meta tag ini, teks dan tombol akan mengecil seperti semut!",
            "Meta tag viewport memerintahkan smartphone menyamakan skala layar 1:1.",
            "Atribut wajib: width=device-width dan initial-scale=1.0.",
            "Merupakan syarat mutlak seluruh aplikasi mobile hybrid modern."
        ],
        code_snippet="""<!-- Baris Sakti yang Wajib Ada di Setiap Aplikasi Mobile -->
<meta name="viewport" 
      content="width=device-width, initial-scale=1.0, maximum-scale=1.0, user-scalable=no">

<!-- Penjelasan Parameter: -->
<!-- width=device-width  : Samakan lebar web dengan lebar fisik HP -->
<!-- initial-scale=1.0   : Tampilkan skala 100% tanpa zoom out -->
<!-- user-scalable=no    : Kunci skala agar tidak melar saat dicubit -->""",
        filename="slide_10_viewport_meta_scaling.html",
        tip="File contoh ini memiliki tombol untuk mematikan viewport dan melihat efek zoom out!",
        tag_color=COLOR_MINT,
        full_code_file="slide_10_viewport_meta_scaling.html"
    )

    # Slide 11: 5 Perintah Sakti Ionic CLI
    builder.add_concept_with_code(
        tag="IONIC CLI",
        title="5 Perintah Sakti Ionic CLI untuk Pengembangan Proyek Penuh",
        bullets=[
            "ionic start: Mengunduh kerangka kerja awal bertema Ionic Vue 3.",
            "cd <nama-folder>: Berpindah masuk ke direktori proyek sebelum mengetik perintah.",
            "ionic serve: Menjalankan server lokal dengan Live Reload di localhost:8100.",
            "ionic build: Mengemas seluruh kode Vue menjadi aset web teroptimasi (dist/).",
            "npx cap sync android: Mengirim hasil build web ke dalam proyek Android native."
        ],
        code_snippet="""# 1. Buat proyek baru bertema Ionic Vue:
ionic start myut-mobile blank --type=vue

# 2. Masuk ke folder proyek:
cd myut-mobile

# 3. Jalankan server lokal:
ionic serve

# 4. Hentikan server kapan saja dengan: Ctrl + C""",
        filename="slide_11_ionic_cli_simulator.js",
        tip="Gunakan terminal VS Code terintegrasi agar tidak perlu berpindah-pindah jendela.",
        tag_color=COLOR_YELLOW,
        full_code_file="slide_11_ionic_cli_simulator.js"
    )

    # Slide 12: Struktur Folder Proyek Standar
    builder.add_concept_with_code(
        tag="STRUKTUR FOLDER",
        title="Anatomi Folder Proyek Ionic Vue: Tahu Persis di Mana Menulis Kode",
        bullets=[
            "src/views/: Tempat utama mahasiswa membuat tampilan layar aplikasi baru.",
            "src/theme/variables.css: Tempat mengganti palet warna biru UT dan Dark Mode.",
            "src/router/index.ts: Tempat mendaftarkan rute navigasi antar-halaman.",
            "capacitor.config.json: Berkas identitas aplikasi (App ID & Nama Aplikasi).",
            "public/: Berkas aset statis seperti logo kampus UT dan favicon."
        ],
        code_snippet="""myut-app/
├── 📁 public/               # Aset gambar & ikon logo kampus
├── 📁 src/                  # JANTUNG KODE MAHASISWA:
│   ├── 📁 views/            # Halaman aplikasi (Home.vue, Profil.vue)
│   ├── 📁 theme/            # variables.css (Warna Tema & Dark Mode)
│   ├── 📁 router/           # index.ts (Alur Navigasi Antar Halaman)
│   └── 📄 App.vue           # Komponen akar pembungkus
└── 📄 capacitor.config.json # Konfigurasi ID Android & Nama Paket""",
        filename="slide_12_struktur_folder_proyek.js",
        tip="Jangan pernah mengedit berkas di dalam folder node_modules secara manual!",
        tag_color=COLOR_CYAN,
        full_code_file="slide_12_struktur_folder_proyek.js"
    )

    # Slide 13: USB Debugging & scrcpy
    builder.add_concept_with_code(
        tag="HEMAT RAM",
        title="Pengujian HP Fisik Hemat RAM via USB Debugging & scrcpy",
        bullets=[
            "Solusi penyelamat laptop mahasiswa: hemat RAM (konsumsi < 70MB).",
            "Langkah 1: Aktifkan 'Opsi Pengembang' dengan mengetuk Build Number 7 kali.",
            "Langkah 2: Aktifkan sakelar 'Proses Debug USB (USB Debugging)'.",
            "Langkah 3: Sambungkan kabel data USB dari smartphone ke laptop.",
            "Langkah 4: Jalankan alat gratis scrcpy untuk menampilkan layar HP di PC."
        ],
        code_snippet="""# Alur Pengujian Smartphone Fisik:
1. Hubungkan ponsel via kabel USB
2. Buka aplikasi scrcpy di laptop:
   scrcpy.exe

3. Buka peramban di ponsel lalu ketik alamat IP laptop:
   http://192.168.1.XX:8100
   
*Hasil: Aplikasi berjalan langsung di hardware HP asli 60 FPS!""",
        filename="slide_13_usb_debugging_scrcpy_guide.md",
        tip="Pastikan kabel USB mendukung transfer data, bukan kabel charger daya biasa.",
        tag_color=COLOR_MINT,
        full_code_file="slide_13_usb_debugging_scrcpy_guide.md"
    )

    # Slide 14: Hot Module Replacement (HMR)
    builder.add_concept_with_code(
        tag="LIVE RELOAD",
        title="Bagaimana Hot Module Replacement (HMR) Membuat Koding Menyenangkan?",
        bullets=[
            "Tradisional: Ubah kode -> Kompilasi 2 menit -> Pasang ulang APK -> Uji coba.",
            "Modern (Ionic + Vite HMR): Ubah kode -> Tekan Ctrl+S -> Layar terupdate < 0.5 detik!",
            "Status data di formulir tidak terhapus saat pembaruan komponen terjadi.",
            "Sangat mempercepat proses belajar dan eksperimen logika antarmuka.",
            "Menghilangkan rasa bosan menunggu proses kompilasi yang lama."
        ],
        code_snippet="""// Logika Pembaruan Instan Vite Engine:
// 1. Vite mendeteksi berkas file.vue mengalami perubahan
// 2. Mengirim sinyal WebSocket ke browser: 'update component'
// 3. Browser mengganti modul yang berubah tanpa me-refresh halaman!

// Hasil: Angka counter formulir tetap utuh, 
// sementara warna latar belakang berubah seketika!""",
        filename="slide_14_hmr_live_reload_demo.html",
        tip="Simpan perubahan kode dengan Ctrl+S secara berkala untuk memicu Live Reload.",
        tag_color=COLOR_YELLOW,
        full_code_file="slide_14_hmr_live_reload_demo.html"
    )

    # Slide 15: Troubleshooting PowerShell
    builder.add_concept_with_code(
        tag="FIRST-AID ERROR",
        title="Pertolongan Pertama: Mengatasi Galat Kebijakan PowerShell Windows",
        bullets=[
            "Pesan Error: 'File ... cannot be loaded because running scripts is disabled'.",
            "Penyebab: Proteksi default keamanan sistem operasi Windows terhadap skrip otomatis.",
            "Solusi Cepat: Buka terminal PowerShell lalu jalankan perintah Set-ExecutionPolicy.",
            "Alternatif Tanpa Pengaturan: Gunakan terminal Git Bash di VS Code.",
            "Ingat: Error ini bukan tanda kegagalan atau kerusakan pada laptop Anda!"
        ],
        code_snippet="""# Solusi 1 Langkah Tuntas di Terminal PowerShell:
Set-ExecutionPolicy -Scope CurrentUser -ExecutionPolicy RemoteSigned

# Tekan Enter, lalu ketik huruf 'Y' untuk konfirmasi.
# Setelah itu coba ketik kembali:
ionic -v

# Selesai! Versi Ionic CLI akan tampil normal.""",
        filename="slide_15_troubleshooting_powershell_cli.md",
        tip="Jika ragu mengubah izin Windows, cukup pilih Git Bash sebagai terminal default Anda.",
        tag_color=COLOR_CORAL,
        full_code_file="slide_15_troubleshooting_powershell_cli.md"
    )

    # Slide 16: Cheatsheet Terminal & Git
    builder.add_concept_with_code(
        tag="CHEATSHEET",
        title="Cheatsheet Perintah Terminal & Git untuk Tugas Kuliah Mahasiswa UT",
        bullets=[
            "Navigasi: cd (pindah folder), cd .. (naik 1 tingkat), dir/ls (lihat isi folder).",
            "Git Status: Memeriksa berkas apa saja yang baru selesai Anda koding.",
            "Git Add & Commit: Menyimpan catatan riwayat pengerjaan tugas mandiri.",
            "Git Push: Mengirimkan berkas ke akun GitHub pribadi untuk dinilai tutor.",
            "Gunakan tombol TAB pada keyboard untuk auto-complete nama folder tanpa typo!"
        ],
        code_snippet="""# Alur Standar Pengumpulan Tugas Praktik di Git:
git status                  # 1. Cek perubahan berkas
git add .                   # 2. Tandai seluruh berkas
git commit -m "tugas 1"     # 3. Simpan riwayat revisi
git push origin main        # 4. Kirim ke repositori online

# Tips Pro: Ketik 'cd con' lalu tekan tombol TAB!""",
        filename="slide_16_cheatsheet_terminal_git.md",
        tip="Biasakan membuat commit dengan pesan yang jelas agar memudahkan pelacakan bug.",
        tag_color=COLOR_PURPLE,
        full_code_file="slide_16_cheatsheet_terminal_git.md"
    )

    # Slide 17: Lab Quest 01
    builder.add_lab_quest(
        meeting_num=1,
        title="Membangun Kartu Tanda Mahasiswa (KTM) Digital Interaktif",
        goals=[
            "Gunakan Web Components resmi Ionic via CDN (@ionic/core).",
            "Tampilkan identitas mahasiswa UT: Nama, NIM 9 digit, Program Studi, & Fakultas FST.",
            "Terapkan palet warna resmi biru Universitas Terbuka (#005691).",
            "Sediakan tombol interaktif 'Verifikasi Keaslian Kartu' dengan umpan balik visual."
        ],
        success_criteria=[
            "Berkas HTML mandiri dapat dibuka langsung di Google Chrome (Zero Installation).",
            "Tampilan adaptif dan proporsional saat diuji via Device Toolbar (Ctrl+Shift+M).",
            "Komponen ion-card, ion-list, dan ion-button tersusun rapi tanpa galat konsol.",
            "Tombol verifikasi berhasil menampilkan status aktif mahasiswa saat disentuh."
        ],
        full_code_file="slide_17_lab_quest_01_profil_mahasiswa.html"
    )

    # Slide 18: Preview Sesi 02
    builder.add_concept_with_code(
        tag="PREVIEW SESI 02",
        title="Menatap Sesi 02: Beralih ke Reaktivitas Modern Vue.js 3",
        bullets=[
            "Hari ini: Mengubah teks antarmuka dengan cara kuno (document.getElementById).",
            "Kelemahan cara kuno: Melelahkan dan rawan salah ketik ID jika aplikasi bertambah besar.",
            "Pekan depan: Mempelajari sistem reaktivitas modern Vue 3 (ref dan reactive).",
            "Keajaiban Vue: Data di variabel berubah, layar antarmuka otomatis terupdate!",
            "Sampai jumpa di Sesi 02: Menguasai Reaktivitas Frontend Modern!"
        ],
        code_snippet="""// Perbandingan Paradigma Kuno vs Reaktif Modern:
// Cara Kuno (Imperatif):
const el = document.getElementById('status');
el.innerText = 'Nilai Baru'; // Rawan typo ID!

// Cara Modern Vue 3 (Reaktif Deklaratif):
const status = ref('Nilai Baru');
// Di template HTML cukup tulis: {{ status }}
// Nilai di layar berubah otomatis seketika data diganti!""",
        filename="slide_18_preview_sesi_02_vue.html",
        tip="Pelajari Modul 2 BMP STSI4303 sebelum menghadiri pertemuan Sesi 02 pekan depan.",
        tag_color=COLOR_YELLOW,
        full_code_file="slide_18_preview_sesi_02_vue.html"
    )

    # Simpan berkas PPTX
    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_file = os.path.join(out_dir, "SESI_01_Pengantar_dan_Lingkungan_Ionic.pptx")
    builder.save(out_file)


if __name__ == "__main__":
    build_sesi_01_deck()
